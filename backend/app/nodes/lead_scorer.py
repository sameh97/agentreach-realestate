"""
Nodes 5 & 6 — Lead Scorer + CSV/XLSX Delivery (Real Estate vertical)
=====================================================================
Scoring here is tuned for who actually BUYS this list: vendors selling
services or products INTO real estate (coaching, E&O insurance,
marketing/CRM tools, title/escrow partners, loan officer referral
partnerships). For that buyer, an agent/brokerage is a *better* lead
when it looks established, active, sizeable, and reachable — not
necessarily when it has the highest star rating (that matters far less
to a B2B vendor than it would to a consumer choosing an agent).

Score breakdown (100 pts):
  Team / brokerage size ............ 22 pts  (bigger org = bigger deal size)
  Established (years in business) .. 15 pts
  Digital maturity (IDX + site) .... 15 pts  (active, maintained business)
  Email deliverability ............. 25 pts  (can you actually reach them)
  Specialization clarity ........... 10 pts  (real, positioned practice)
  Social proof (testimonials) ...... 8 pts
  License / compliance detected .... 5 pts
"""
import os, io, csv, logging
from datetime import datetime
from pathlib import Path
from app.agents.state import LeadState

logger = logging.getLogger(__name__)
OUTPUT_DIR = Path(os.getenv("OUTPUT_DIR", "/tmp/leads"))
OUTPUT_DIR.mkdir(parents=True, exist_ok=True)

COLS = [
    "score", "name", "primary_email", "email_verified", "email_catchall",
    "email_status", "team_size", "specializations", "years_in_business",
    "license_detected", "idx_detected", "testimonial_count",
    "service_area_count", "phone", "website", "address",
    "rating", "review_count", "category", "source",
]

_TEAM_SIZE_PTS = {"brokerage": 22, "large_team": 17, "small_team": 11, "solo": 6, "unknown": 3}


# ── Scorer ────────────────────────────────────────────────────────────────────

def _score(lead: dict) -> int:
    s = 0

    # Team / brokerage size (22 pts) — bigger org, bigger deal for a B2B vendor
    s += _TEAM_SIZE_PTS.get(lead.get("team_size", "unknown"), 3)

    # Established (15 pts)
    years = int(lead.get("years_in_business") or 0)
    s += 15 if years >= 10 else 10 if years >= 5 else 5 if years >= 2 else 0

    # Digital maturity (15 pts) — active, maintained web presence
    if lead.get("idx_detected"):   s += 10
    if lead.get("website"):        s += 3
    if lead.get("recent_activity"): s += 2

    # Email deliverability (25 pts) — can you actually reach them
    if lead.get("email_verified"):    s += 25
    elif lead.get("email_catchall"):  s += 15
    elif lead.get("email_valid"):     s += 8

    # Specialization clarity (10 pts) — a positioned practice, not a placeholder
    specs = lead.get("specializations") or []
    s += 10 if len(specs) >= 2 else 6 if len(specs) == 1 else 0

    # Social proof (8 pts)
    testi = int(lead.get("testimonial_count") or 0)
    s += 8 if testi >= 5 else 4 if testi >= 1 else 0

    # License / compliance detected (5 pts)
    if lead.get("license_detected"): s += 5

    return min(100, s)


def score_leads_node(state: LeadState) -> LeadState:
    leads = state.get("verified_leads", [])
    scored = sorted(
        [{**l, "score": _score(l)} for l in leads],
        key=lambda l: l["score"], reverse=True,
    )
    high = sum(1 for l in scored if l["score"] >= 70)
    logger.info(f"[scorer] {len(scored)} leads — {high} high-quality (≥70)")
    return {**state, "scored_leads": scored}


# ── Delivery ──────────────────────────────────────────────────────────────────

def _flatten_for_export(lead: dict) -> dict:
    """CSV/XLSX writers choke on list values — join specializations to a string."""
    out = dict(lead)
    specs = out.get("specializations") or []
    out["specializations"] = ", ".join(s.replace("_", " ").title() for s in specs)
    return out


def leads_to_csv_bytes(scored_leads: list[dict]) -> bytes:
    leads = [_flatten_for_export(l) for l in scored_leads]
    buf = io.StringIO()
    w = csv.DictWriter(buf, fieldnames=COLS, extrasaction="ignore")
    w.writeheader()
    w.writerows(leads)
    return buf.getvalue().encode("utf-8")


def leads_to_xlsx_bytes(scored_leads: list[dict]) -> bytes:
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter

    leads = [_flatten_for_export(l) for l in scored_leads]

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Leads"

    hdr_fill = PatternFill("solid", fgColor="0F172A")
    hdr_font = Font(bold=True, color="FFFFFF", size=10)
    for ci, col in enumerate(COLS, 1):
        cell = ws.cell(row=1, column=ci, value=col.replace("_", " ").title())
        cell.fill = hdr_fill
        cell.font = hdr_font
        cell.alignment = Alignment(horizontal="center")

    g = PatternFill("solid", fgColor="DCFCE7")
    y = PatternFill("solid", fgColor="FEF9C3")
    r = PatternFill("solid", fgColor="FEE2E2")

    for ri, lead in enumerate(leads, 2):
        score = lead.get("score", 0)
        fill = g if score >= 70 else y if score >= 40 else r
        for ci, col in enumerate(COLS, 1):
            cell = ws.cell(row=ri, column=ci, value=lead.get(col, ""))
            cell.fill = fill

    for ci, col in enumerate(COLS, 1):
        maxw = max(len(col), *(len(str(l.get(col, ""))) for l in leads[:50])) + 2
        ws.column_dimensions[get_column_letter(ci)].width = min(maxw, 40)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def export_filename(location: str, created_at: datetime, ext: str) -> str:
    loc = (location or "leads")[:40].replace(" ", "_").replace(",", "").replace("/", "-")
    return f"agentreach_{loc}_{created_at.strftime('%Y%m%d_%H%M%S')}.{ext}"


def deliver_leads_node(state: LeadState) -> LeadState:
    leads = state.get("scored_leads", [])
    now   = datetime.now()

    # CSV
    csv_path = OUTPUT_DIR / export_filename(state.get("location", ""), now, "csv")
    csv_path.write_bytes(leads_to_csv_bytes(leads))
    logger.info(f"[deliver] CSV → {csv_path}  ({len(leads)} rows)")

    # XLSX
    xlsx_path = None
    try:
        xlsx_path = OUTPUT_DIR / export_filename(state.get("location", ""), now, "xlsx")
        xlsx_path.write_bytes(leads_to_xlsx_bytes(leads))
        logger.info(f"[deliver] XLSX → {xlsx_path}")
    except ImportError:
        xlsx_path = None
        logger.warning("[deliver] openpyxl not installed, skipping XLSX")
    except Exception as e:
        xlsx_path = None
        logger.error(f"[deliver] XLSX error: {e}")

    return {
        **state,
        "final_csv_path":  str(csv_path),
        "final_xlsx_path": str(xlsx_path) if xlsx_path else "",
        "status": "done",
    }
